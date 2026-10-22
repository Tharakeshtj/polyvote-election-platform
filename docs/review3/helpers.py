"""Document helpers that reproduce the Review I/II formatting (Times New Roman 12, justified body,
bold numbered headings, italic figure captions below figures, bold table captions below tables)."""
import copy
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

FONT = 'Times New Roman'
MONO = 'Consolas'


class Ctx:
    """Holds the document plus figure/table counters. Run the build twice so cross-references resolve."""

    def __init__(self, template: Path, prev_labels=None):
        self.doc = docx.Document(str(template))
        self.labels = {}
        self.prev = prev_labels or {}
        self.fig = 0
        self.tab = 0
        body = self.doc.element.body
        bullet = next(p for p in self.doc.paragraphs if p.style.name == 'List Paragraph')
        self.bullet_ppr = copy.deepcopy(bullet._p.pPr)
        for el in list(body):
            if el.tag != qn('w:sectPr'):
                body.remove(el)

    # ------------------------------------------------------------- text
    def _runs(self, p, text, size=12, base_bold=False, base_italic=False, color=None):
        """Inline markup: **bold**, *italic*, `code`."""
        for tok in re.split(r'(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)', text):
            if not tok:
                continue
            bold, italic, mono = base_bold, base_italic, False
            if tok.startswith('**') and tok.endswith('**'):
                tok, bold = tok[2:-2], True
            elif tok.startswith('`') and tok.endswith('`'):
                tok, mono = tok[1:-1], True
            elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
                tok, italic = tok[1:-1], True
            r = p.add_run(tok)
            r.font.name = MONO if mono else FONT
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO if mono else FONT)
            r.font.size = Pt(size - 1.5 if mono else size)
            r.bold = bold or None
            r.italic = italic or None
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def p(self, text, align='justify', before=0, after=8, size=12, bold=False, italic=False, keep_next=False):
        para = self.doc.add_paragraph()
        para.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
        pf = para.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = keep_next or None
        self._runs(para, text, size, bold, italic)
        return para

    def h1(self, text):
        para = self.p(text, align='left', before=14, after=8, bold=True, keep_next=True)
        return para

    def h2(self, text):
        return self.p(text, align='left', before=10, after=6, bold=True, italic=True, keep_next=True)

    def h3(self, text):
        return self.p(text, align='left', before=6, after=4, italic=True, keep_next=True)

    def bullet(self, text, size=12):
        para = self.doc.add_paragraph(style='List Paragraph')
        para._p.remove(para._p.pPr)
        para._p.insert(0, copy.deepcopy(self.bullet_ppr))
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_after = Pt(4)
        self._runs(para, text, size)
        return para

    def code(self, text, size=8.5):
        lines = text.strip('\n').split('\n')
        for i, line in enumerate(lines):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Emu(127000)
            pf.space_before = Pt(0)
            pf.space_after = Pt(10) if i == len(lines) - 1 else Pt(0)
            pf.keep_together = True
            if i < len(lines) - 1:
                pf.keep_with_next = True
            r = para.add_run(line)
            r.font.name = MONO
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            r.font.size = Pt(size)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------- floats
    def ref(self, label):
        return self.prev.get(label, '?')

    def figure(self, path, caption, label, width=6.3):
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(path)
        self.fig += 1
        self.labels[label] = f'Figure {self.fig}'
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before, para.paragraph_format.space_after = Pt(6), Pt(3)
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(11)
        self._runs(cap, f'Figure {self.fig}: {caption}', size=10.5, base_italic=True)

    def figure_row(self, paths, caption, label, width_each=3.0):
        """Several images side by side under one caption."""
        for p in paths:
            if not Path(p).exists():
                raise FileNotFoundError(p)
        self.fig += 1
        self.labels[label] = f'Figure {self.fig}'
        t = self.doc.add_table(rows=1, cols=len(paths))
        t.alignment = 1
        for cell, p in zip(t.rows[0].cells, paths):
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(p), width=Inches(width_each))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(4), Pt(11)
        self._runs(cap, f'Figure {self.fig}: {caption}', size=10.5, base_italic=True)

    def table(self, headers, rows, widths, caption, label, size=9.5, align=None):
        self.tab += 1
        self.labels[label] = f'Table {self.tab}'
        total = 9360
        w = [int(total * x / sum(widths)) for x in widths]
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        tblPr = t._tbl.tblPr
        for el in list(tblPr):
            tblPr.remove(el)
        tw = OxmlElement('w:tblW'); tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(total)); tblPr.append(tw)
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single'); b.set(qn('w:color'), 'auto'); b.set(qn('w:sz'), '4')
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
        for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), w):
            gc.set(qn('w:w'), str(wd))
        for ri, row in enumerate([headers] + rows):
            tr = t.rows[ri]
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                hdr = OxmlElement('w:tblHeader'); trPr.append(hdr)
            cant = OxmlElement('w:cantSplit'); tr._tr.get_or_add_trPr().append(cant)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                tcPr = cell._tc.get_or_add_tcPr()
                tcw = OxmlElement('w:tcW'); tcw.set(qn('w:type'), 'dxa'); tcw.set(qn('w:w'), str(w[ci])); tcPr.append(tcw)
                tb = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    b = OxmlElement(f'w:{side}'); b.set(qn('w:val'), 'single'); b.set(qn('w:color'), 'BFBFBF'); b.set(qn('w:sz'), '4'); tb.append(b)
                tcPr.append(tb)
                if ri == 0:
                    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F2F2F2'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:val'), 'clear'); tcPr.append(shd)
                mar = OxmlElement('w:tcMar')
                for side, v in (('top', 60), ('left', 90), ('bottom', 60), ('right', 90)):
                    m = OxmlElement(f'w:{side}'); m.set(qn('w:type'), 'dxa'); m.set(qn('w:w'), str(v)); mar.append(m)
                tcPr.append(mar)
                para = cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_after = Pt(0)
                self._runs(para, str(val), size=size, base_bold=(ri == 0))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(4), Pt(10)
        self._runs(cap, f'Table {self.tab}: {caption}', size=11, base_bold=True)

    # ------------------------------------------------------------- footer
    def page_numbers(self):
        sec = self.doc.sections[0]
        para = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for child in list(para._p):
            if child.tag != qn('w:pPr'):
                para._p.remove(child)
        run = para.add_run()
        run.font.name = FONT
        run.font.size = Pt(10)
        for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
            if kind:
                fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
            else:
                it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = text; run._r.append(it)

    def save(self, path):
        self.doc.save(str(path))
