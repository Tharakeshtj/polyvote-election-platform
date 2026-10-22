"""Build docs/PolyVote_IEEE_Paper.docx: a two-column IEEE-style conference paper from the same data as the report.

    python docs/review3/paper.py
"""
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import docx                                              # noqa: E402
from docx.enum.section import WD_SECTION                 # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH            # noqa: E402
from docx.oxml import OxmlElement                        # noqa: E402
from docx.oxml.ns import qn                              # noqa: E402
from docx.shared import Inches, Pt                       # noqa: E402

import data                                              # noqa: E402

S = data.load()
ROOT = data.ROOT
OUT = ROOT / 'docs' / 'PolyVote_IEEE_Paper.docx'
FONT = 'Times New Roman'

ids = [r['id'] for r in json.loads((HERE / 'references.json').read_text(encoding='utf-8'))]
ieee = [re.sub(r'^\[\d+\]\s*', '', ln) for ln in (HERE / 'references_ieee.txt').read_text(encoding='utf-8').split('\n') if ln.strip()]
assert len(ids) == len(ieee)
TEXT = dict(zip(ids, ieee))
cited = []


def cite(m):
    keys = [k.strip() for k in m.group(1).split(',')]
    nums = []
    for k in keys:
        if k not in TEXT:
            raise KeyError(k)
        if k not in cited:
            cited.append(k)
        nums.append(cited.index(k) + 1)
    return '[' + ', '.join(str(n) for n in sorted(nums)) + ']'


doc = docx.Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
sec.left_margin = sec.right_margin = Inches(0.625)
st = doc.styles['Normal']
st.font.name, st.font.size = FONT, Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def runs(p, text, size=10, bold=False, italic=False):
    text = re.sub(r'\{([^}]+)\}', cite, text)
    for tok in re.split(r'(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)', text):
        if not tok:
            continue
        b, i, mono = bold, italic, False
        if tok.startswith('**'):
            tok, b = tok[2:-2], True
        elif tok.startswith('`'):
            tok, mono = tok[1:-1], True
        elif tok.startswith('*') and len(tok) > 2:
            tok, i = tok[1:-1], True
        r = p.add_run(tok)
        r.font.name = 'Consolas' if mono else FONT
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas' if mono else FONT)
        r.font.size = Pt(size - 1 if mono else size)
        r.bold, r.italic = b or None, i or None


def para(text, align='justify', size=10, before=0, after=4, bold=False, italic=False, indent=True, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    if indent and align == 'justify':
        pf.first_line_indent = Inches(0.17)
    pf.keep_with_next = keep_next or None
    runs(p, text, size, bold, italic)
    return p


def h1(text):
    p = para(text.upper(), 'center', 10, before=8, after=4, indent=False, keep_next=True)
    for r in p.runs:
        r.font.small_caps = False
    return p


def h2(text):
    return para(text, 'left', 10, before=4, after=2, italic=True, indent=False, keep_next=True)


def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Inches(0.2), Inches(-0.15)
    pf.space_after, pf.line_spacing = Pt(2), 1.0
    runs(p, '• ' + text, 10)


def set_columns(section, n, space=0.25):
    sectPr = section._sectPr
    for el in sectPr.findall(qn('w:cols')):
        sectPr.remove(el)
    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), str(int(space * 1440)))
    sectPr.append(cols)


def figure(path, caption, width=3.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    para(caption, 'center', 8, after=6, indent=False)


def table(headers, rows, widths, caption, size=8):
    para(caption, 'center', 8, before=4, after=2, indent=False, keep_next=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = 1
    for ri, row in enumerate([headers] + rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.width = Inches(widths[ci])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            runs(p, str(val), size, bold=(ri == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


LABEL = {'detected': 'Detected', 'detected_manual_only': 'Counters only', 'undetected': 'Not detected'}
base, pat, mg = S['bench']['baseline'], S['bench']['patched'], S['mongo']
tam = S['tamper']
nd = lambda k: sum(1 for t in tam if t[k] == 'detected')  # noqa: E731

# ------------------------------------------------------------------ title block (single column)
para('PolyVote: Database-Enforced Integrity and a Keyed Audit Ledger for Small Ranked-Choice Elections', 'center', 22, after=8, indent=False)
para('Tharakesharasu Thennarasu, M Pavan Durga Karthik and Shivam Malge', 'center', 11, after=0, indent=False)
para('School of Computer Science and Engineering, Course BACSE202 Database Systems', 'center', 10, after=0, indent=False, italic=True)
para('Faculty guide: Dr. Katari Balakrishna', 'center', 10, after=8, indent=False, italic=True)

abstract = (
    '**Abstract**—Small elections are usually run on a general-purpose database that an administrator can edit without leaving a trace. '
    'Cryptographic end-to-end verifiable systems {adida2008, cortier2019} address this but carry operational burdens that campus and club elections rarely accept. '
    'We present PolyVote, a multi-tenant ranked-choice (instant-runoff) election platform that raises the cost of undetected tampering using database features alone. '
    'Constraints and a trigger enforce one vote per voter; voters and anonymous ballots are stored without a linking column; each ballot is sealed in an HMAC-keyed, '
    'hash-chained ledger kept in MongoDB Atlas; and an integrity sweep reconciles ballots, receipts, voter flags and the sealed candidate roster. A standalone '
    'email one-time-code module adds a second authentication factor. We evaluate the system with an executable suite of 47 scenarios. Against the earlier design, rogue-'
    f'administrator scenarios detected automatically rise from {nd("baseline")} to {nd("patched")} of {len(tam)}, including forged receipts and consistent chain rewrites that a plain hash chain cannot stop; '
    'sweep time rises by about 14% at 10,000 ballots. One scenario, roster injection, remains undetected, and we state the limits of a design that is not end-to-end verifiable.')
p = para(abstract, 'justify', 9, after=4, indent=False, bold=False)
for r in p.runs:
    r.italic = False
para('**Index Terms**—electronic voting, instant-runoff voting, tamper-evident logging, HMAC, database integrity, polyvalent persistence, one-time password.'.replace('polyvalent', 'polyglot'),
     'justify', 9, after=6, indent=False)

new = doc.add_section(WD_SECTION.CONTINUOUS)
set_columns(new, 2)

# ------------------------------------------------------------------ I
h1('I. Introduction')
para('Student councils, clubs and societies run elections on spreadsheets, survey forms or small web applications backed by one database. Whoever administers that database can '
     'change the outcome, and nothing in the system can show that a result was altered. Research on electronic voting has focused on cryptographic protocols that avoid '
     'trusting the server {adida2008, cortier2019, clarkson2008, ryan2009, chaum2009} and on the failures of deployed internet voting {springall2014, specter2020, specter2021, nasem2018}. '
     'Between the two lies a large population of small elections for which verifiable cryptographic voting is impractical, yet where much more can be done by using the database properly.')
para('PolyVote treats the database as the place where election rules are enforced and adds an independent append-only audit trail that can be checked against it. '
     'Its contributions are: (1) a database-centric design that enforces one-vote and ballot-validity rules in the engine and separates voters from ballots; '
     '(2) an HMAC-keyed hash-chained ledger in a document store with a sealed candidate roster; (3) an extended integrity sweep; (4) a standalone email one-time-code '
     'module; and (5) a reproducible evaluation of 47 scenarios against both the earlier design and the final one, with an explicit statement of what is not claimed.')

# ------------------------------------------------------------------ II
h1('II. Related Work')
para('Helios {adida2008} and Belenios {cortier2019} provide verifiable tallies using encrypted ballots; Civitas {clarkson2008} adds coercion resistance; Prêt à Voter {ryan2009} and Scantegrity II {chaum2009} '
     'target paper-based verification. Analyses of deployed systems found significant weaknesses in Estonian internet voting {springall2014} and in mobile and web products {specter2020, specter2021}, and a National Academies '
     'study concluded internet voting could not currently be secured {nasem2018}. Blockchain voting compounds these problems {park2021}.')
para('Hash chaining for time-stamping {haber1991} and secure audit logs {schneier1999} underlie tamper-evident logging {crosby2009, rfc6962}; database work covers tamper detection in audit logs {snodgrass2004}, '
     'forensic analysis of tampering {pavlou2008} and ledger tables in a commercial engine {antonopoulos2021}. Instant-runoff outcomes can hinge on few ballots, motivating margin computation and '
     'risk-limiting audits {magrino2011, blom2019}. Multi-tenant schema mapping {aulbach2008}, polyglot persistence {sadalage2012} and database security {bertino2005} complete the background. '
     'For authentication, NIST SP 800-63B separates memorised secrets from possession factors {nist80063b4}; email {garfinkel2003} and mobile channels {dmitrienko2014} are convenient but only as strong as the mailbox or device.')

# ------------------------------------------------------------------ III
h1('III. Threat Model')
para('We consider seven adversaries: a voter with a leaked password (A1); a network observer (A2, out of scope beyond TLS); a rogue administrator of the election database (A3); an administrator of the ledger store (A4); '
     'an insider with both stores but no secrets (A5); an insider who also holds the key (A6); and roster injection through the administrative path (A7). '
     'Goals are: no double vote; only rostered, verified voters vote; ballots anonymous in storage; malformed ballots rejected; modification, deletion or insertion detectable; the ledger not silently rewritable without the key; '
     'candidate attribution fixed after voting; and voters able to confirm their ballot is recorded. A6 and A7 are not covered.')

figure(data.fig('threat_model.png'), 'Fig. 1. Threat model: adversaries, capabilities and how each is treated.', 3.45)
h2('Assumptions')
para('The application host and its environment secrets (database connection string, ledger key, SMTP credentials) are not controlled by the adversary except in A6. A voter controls their mailbox. '
     'The election officer registers the roster honestly. Coercion, vote selling, denial of service and compromise of the voter\'s own device are out of scope.')

# ------------------------------------------------------------------ IV
h1('IV. Design')
para('PolyVote is one FastAPI service with a React front end. It uses three stores separated by what each is trusted with: a relational election database, a separate verification database for emails and one-time codes, '
     'and a MongoDB Atlas ledger, with application secrets held outside all of them (Fig. 2). An administrator of one store cannot touch the others, and only the application host sees all three, '
     'which is why host compromise is outside the model.')
figure(data.fig('architecture.png'), 'Fig. 2. Architecture: election database, separate verification database and MongoDB ledger, with secrets outside all stores.', 3.45)
h2('A. Data model and rules')
para('One shared schema serves many organisations {aulbach2008}. Six tables hold organisations, elections, candidates, the roster and ranked ballot rows. The engine enforces: a trigger that aborts a second '
     '`has_voted` update, CHECK constraints on flags and rank (1–3), a composite key that prevents ranking a candidate twice, per-election uniqueness of registration numbers, and foreign keys {pg_constraints, pg_triggers}. '
     'A ballot is stored under a random token; no column links it to a voter, and voter email addresses live in a different database (Fig. 3). An automated check confirms the two ballot and voter tables share no column other than the election identifier.')
figure(data.fig('er.png'), 'Fig. 3. Entity-relationship diagram: election database, verification database and ledger block.', 3.45)
h2('B. Casting a ballot')
para('Casting is serialised and transactional. After authentication the server takes a write lock, re-checks the election status, validates the ranking and confirms the voter has not voted. In one transaction it sets the voted flag '
     'and inserts the ranked rows under a new random token, then appends a sealed block to the ledger and invalidates the session. The append follows the commit, so a failure between them leaves a ballot with no receipt, '
     'which the sweep reports as injected; we chose the variant that fails visibly. Tallying is instant-runoff: first preferences are counted among continuing candidates, a candidate with more than half wins, otherwise the weakest is eliminated '
     'and ballots transfer. Ties for last place are broken by candidate order, and the interface flags such rounds.')
h2('C. Keyed audit ledger')
para('Each ballot yields one block holding its index, election, token, a SHA-256 of the ranked choices, a SHA-256 of the candidate roster, the previous hash and its own hash. The block hash is HMAC-SHA-256 under a secret '
     'held outside both stores, so an attacker who can write the ledger but lacks the key cannot rebuild a consistent chain, which a plain hash chain permits. In MongoDB a unique index on the block position '
     'prevents two writers from forking the chain; a local file is the fallback.')
h2('D. Integrity sweep')
para('The sweep applies five rules: modified ballots (recomputed hash differs from the receipt), deleted ballots (receipt without rows), injected ballots (rows without receipt), changed roster (candidate names differ from the sealed hash), '
     'and count mismatch among voter flags, distinct ballots and receipts. A separate chain check verifies every link and block hash.')
h2('E. Email one-time code')
para('After a correct password the server issues a challenge rather than a session. A six-digit code from a cryptographic generator is stored only as an HMAC, emailed from a background thread, valid for five minutes, '
     'single use, limited to three guesses, with resends limited to one per 30 s and five per challenge, and at most five challenges per voter per 15 minutes. Failure of mail delivery never blocks sign-in; the administrator outbox shows issued codes (Fig. 4). '
     'Voter addresses live in a database separate from ballots, so an email address cannot be joined to a ballot even if the election database is compromised.')
figure(data.fig('otp_sequence.png'), 'Fig. 4. Sign-in with a one-time code.', 3.45)

# ------------------------------------------------------------------ V
h1('V. Evaluation')
para('An executable suite copies the application to a temporary directory, drives the real API, and plays a rogue administrator by editing the database and ledger files directly. '
     'Groups are: 7 database-level prevention tests, 13 API abuse tests, 13 tampering scenarios, 14 OTP tests, and benchmarks at 100, 1,000 and 10,000 ballots (median of three runs). '
     'The baseline is the earlier design, checked out and run unchanged; attacker scripts never receive the key.')
h2('A. Prevention and abuse')
para('Six of seven database-level statements are rejected by the engine. The exception is a dangling candidate reference inserted from a console with default settings, because SQLite enforces foreign keys per connection. '
     'All 13 API abuse tests are refused, including twenty simultaneous casts by one voter, of which exactly one is stored. All 14 OTP tests pass; test F4 initially failed and exposed a defect in which '
     'failed guesses were rolled back with the aborted request, leaving the attempt limit ineffective.')
h2('B. Tampering')
table(['ID', 'Scenario', 'Earlier', 'Final'],
      [[t['id'], t['attack'], LABEL[t['baseline']], LABEL[t['patched']]] for t in tam],
      [0.3, 1.95, 0.5, 0.5], 'TABLE I. Outcome of each tampering scenario.', size=7)
para(f'Under the earlier design {nd("baseline")} of {len(tam)} scenarios are detected automatically, {sum(1 for t in tam if t["baseline"] == "detected_manual_only")} only through counters an officer would have to compare by eye, and '
     f'{sum(1 for t in tam if t["baseline"] == "undetected")} not at all. The final design detects {nd("patched")} and misses one. Forging a receipt (C6) and rebuilding the whole chain (C10) succeed against a plain chain and fail against the keyed one. '
     'Renaming candidates (C11) is caught by the sealed roster. C13, adding a ghost voter and voting as them, is undetected: the ballot is valid, and the system holds no independent record of the roster.')
figure(data.fig('detection.png'), 'Fig. 5. Detection outcome per scenario, earlier design versus final design.', 3.45)
figure(data.fig('admin_sweep_tamper.png'), 'Fig. 6. Sweep after one ballot was deleted and another injected: both are reported although the three counters still read 46/46/46.', 3.45)
para('Fig. 6 matters because it shows why counters alone are not enough: a deletion and an injection cancel out in the counts, yet the sweep names both ballots.')
h2('C. One-time-code tests')
para('The 14 OTP tests check that a correct password alone yields no session; that the right code opens the ballot; that a used code cannot be replayed; that three wrong codes lock the challenge; that expired codes, rapid resends and a sixth '
     'code are refused; that a voter without an email on file cannot proceed; that plaintext codes are absent from the verification database; that repeated sign-ins are rate-limited; that codes are six digits with leading zeros preserved; '
     'that an unreachable mail server does not block sign-in; and that the election database has no email column.')
h2('D. Performance')
table(['Ballots', 'Cast ms', 'Sweep ms', 'Chain ms', 'Tally ms'],
      [[f'{int(n):,}', f"{base[n]['cast_ms']:.0f} / {pat[n]['cast_ms']:.0f}", f"{base[n]['sweep_ms']:.0f} / {pat[n]['sweep_ms']:.0f}",
        f"{base[n]['chain_verify_ms']:.0f} / {pat[n]['chain_verify_ms']:.0f}", f"{base[n]['tally_ms']:.0f} / {pat[n]['tally_ms']:.0f}"] for n in base],
      [0.6, 0.7, 0.7, 0.7, 0.7], 'TABLE II. File-ledger timings, earlier / final design.', size=8)
para(f"Read-side operations grow linearly. The keyed chain and extra rules cost about {(pat['10000']['sweep_ms'] / base['10000']['sweep_ms'] - 1) * 100:.0f}% on the sweep at 10,000 ballots. "
     f"With the file ledger a cast grows from {base['100']['cast_ms']:.0f} ms to {base['10000']['cast_ms']:.0f} ms, since each append rewrites the file. "
     f"Against MongoDB Atlas over a mobile hotspot an append costs about {mg['1000']['append_ms_median']:.0f} ms regardless of chain length; the file is faster at small sizes and the crossover is near ten thousand ballots. "
     f"Verifying the whole chain over the network grows from {mg['100']['verify_ms'] / 1000:.1f} s at 100 blocks to {mg['10000']['verify_ms'] / 1000:.1f} s at 10,000, so a production system should verify incrementally from a trusted checkpoint.")
figure(data.fig('bench.png'), 'Fig. 7. Cast latency and sweep time against ledger size (log scales).', 3.45)
h2('E. Threats to validity')
para('The workload is synthetic and timings come from one laptop and one mobile network, so we report trends rather than absolute values. The attack scenarios are the ones we thought of and are expressed against our own design, '
     'so they are not a complete catalogue. The baseline is our earlier version; we did not run Helios or Belenios, so Section VI is qualitative. Detection also depends on the officer running the checks.')

# ------------------------------------------------------------------ VI
h1('VI. Comparison')
table(['System', 'Assurance', 'Voter verifies', 'Burden'],
      [['Helios, Belenios', 'Encrypted, verifiable tally', 'Yes', 'Moderate'],
       ['Civitas, Prêt à Voter, Scantegrity II', 'Verifiable, some coercion-resistant', 'Yes', 'High'],
       ['Estonian i-voting', 'National crypto system', 'Partial', 'High'],
       ['Commercial mobile/web', 'Vendor trust', 'No', 'Low'],
       ['**PolyVote**', 'Database rules; keyed ledger', 'Recorded, not content', 'Low']],
      [1.3, 1.2, 0.7, 0.5], 'TABLE III. Qualitative comparison (published designs {adida2008, cortier2019, clarkson2008, ryan2009, chaum2009, springall2014, specter2020, specter2021}).', size=7)
table(['Measure', 'Earlier design', 'Final design'],
      [['Tampering detected automatically', f'{nd("baseline")} of {len(tam)}', f'{nd("patched")} of {len(tam)}'],
       ['Forged receipt / chain rewrite', 'Undetected', 'Detected (key)'],
       ['Candidate renaming', 'Undetected', 'Detected (roster seal)'],
       ['DB prevention / API abuse tests', '6/7, 13/13', '6/7, 13/13'],
       ['Voter authentication', 'Password', 'Password + emailed code'],
       ['Ledger storage', 'Local file', 'MongoDB Atlas'],
       ['Sweep at 10,000 ballots', f"{base['10000']['sweep_ms']:.0f} ms", f"{pat['10000']['sweep_ms']:.0f} ms"]],
      [1.4, 0.9, 1.0], 'TABLE IV. Earlier design versus final design.', size=7.5)
para('PolyVote sits at the low-assurance, low-burden end and is not a substitute for verifiable systems where a hostile administrator is the main concern. '
     'We claim an improvement over the earlier design and over unprotected relational election applications, not over end-to-end verifiable systems.')

# ------------------------------------------------------------------ VII
h1('VII. Limitations and Future Work')
bullet('An insider holding the key and both stores can forge a consistent chain (A6). Publishing the chain head to an external witness, as transparency logs do {rfc6962}, would address this.')
bullet('Roster injection (A7) is not detected; a committed or signed roster is needed.')
bullet('Receipts prove a ballot is recorded and counted, not its content {chaum2004}; coercion is not resisted {juels2005}.')
bullet('Passwords use plain SHA-256; salted, slow hashing is recommended {nist800132, owasp_pwd}. Email is a weak factor {garfinkel2003}; TOTP {rfc6238} is a stronger alternative.')
bullet('Foreign-key enforcement depends on the connection; a server database removes this. Oracle with PL/SQL, roles and an entity for election positions were deferred.')

h1('VIII. Conclusion')
para('Using database features alone, PolyVote enforces one-person-one-vote in the engine, stores voters and ballots without a linking column, seals each ballot in a keyed ledger and exposes tampering through a reconciling sweep, '
     f'with a second authentication factor for sign-in. Of {len(tam)} tampering scenarios, the earlier design detected {nd("baseline")} automatically and the final design detects {nd("patched")}, at a small run-time cost. '
     'We report one undetected scenario and the limits of the approach, and position the system as a practical option for small elections rather than a rival to cryptographic voting.')

# ------------------------------------------------------------------ refs
h1('References')
for i, k in enumerate(cited, start=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(0.3), Inches(-0.3)
    p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(1), 1.0
    runs(p, f'[{i}] ' + TEXT[k].replace('{', '').replace('}', ''), 7.5)

doc.save(str(OUT))
print('wrote', OUT, '| cited', len(cited), 'of', len(ids))
